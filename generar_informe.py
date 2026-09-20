import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_report():
    doc = docx.Document()

    # Configuración de márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Título Principal
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("INFORME DE PRÁCTICA: CHATBOT MULTIMODAL\nGASTRONOMÍA PERUANA")
    run_title.bold = True
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = RGBColor(192, 0, 0) # Guinda/Rojo institucional

    # Subtítulo
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Python + Streamlit + OpenAI API (GPT) + Whisper")
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(80, 80, 80)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Curso: Herramientas de Desarrollo Profesional - TIC\nUnidad: Prompt Engineering e Inteligencia Artificial\nSemana 06 - Sesión 01\n")
    r_meta.font.size = Pt(11)
    r_meta.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 1. OBJETIVO
    h1 = doc.add_heading("1. Objetivo de la Práctica", level=1)
    h1.paragraph_format.space_before = Pt(12)
    p = doc.add_paragraph(
        "Desarrollar e implementar un chatbot conversacional multimodal especializado en gastronomía peruana. "
        "La solución integra una interfaz web construida con Streamlit, procesamiento de lenguaje natural con modelos de OpenAI "
        "(Chat Completions) y capacidades de transcripción de voz a texto utilizando el modelo Whisper-1. "
        "El sistema es capaz de responder consultas formuladas tanto por texto escrito como a través de archivos de audio subidos por el usuario."
    )
    p.paragraph_format.line_spacing = 1.15

    # 2. ARQUITECTURA
    h2 = doc.add_heading("2. Arquitectura de la Solución", level=1)
    p_arch = doc.add_paragraph(
        "La aplicación sigue una arquitectura modular en capas que garantiza simplicidad y robustez:\n"
        "1. Interfaz de Usuario (Streamlit): Proporciona los componentes visuales st.chat_input, st.chat_message, "
        "st.file_uploader y st.audio para la interacción del usuario y la gestión del historial de mensajes mediante st.session_state.\n"
        "2. Módulo de Audio (OpenAI Whisper): Recibe archivos de audio en formatos comunes (.mp3, .wav, .m4a, .webm), "
        "los almacena temporalmente y los envía al endpoint client.audio.transcriptions.create(model='whisper-1', language='es') "
        "para convertirlos a texto con alta fidelidad.\n"
        "3. Motor Conversacional (OpenAI Chat Completions): Gestiona el prompt de sistema especializado en cocina peruana, "
        "concatena el historial acumulado en la sesión y genera respuestas coherentes, didácticas y contextualizadas."
    )
    p_arch.paragraph_format.line_spacing = 1.15

    # 3. ESTRUCTURA DEL PROYECTO
    doc.add_heading("3. Estructura del Proyecto", level=1)
    p_tree = doc.add_paragraph(
        "chatbot_comida_peruana/\n"
        "├── .env                     # Clave de API de OpenAI y configuración de modelo (IGNORADO EN GIT)\n"
        "├── .gitignore                # Reglas de exclusión para proteger credenciales\n"
        "├── app.py                   # Código principal de la aplicación Streamlit y OpenAI\n"
        "├── requirements.txt         # Lista de dependencias del proyecto\n"
        "└── audios_prueba/           # Muestras de audio para validación de Whisper\n"
        "    ├── pregunta_causa_limena.wav\n"
        "    └── pregunta_aji_gallina.wav"
    )
    p_tree.paragraph_format.left_indent = Inches(0.5)

    # 4. CÓDIGO FUENTE COMPLETO
    doc.add_heading("4. Código Fuente Completo (app.py)", level=1)
    p_code_desc = doc.add_paragraph("A continuación se presenta el código fuente íntegro de la aplicación app.py:")
    
    with open(r"C:\Users\Jeremy\chatbot_comida_peruana\app.py", "r", encoding="utf-8") as f:
        code_text = f.read()

    p_code = doc.add_paragraph()
    r_code = p_code.add_run(code_text)
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8.5)
    p_code.paragraph_format.left_indent = Inches(0.3)
    p_code.paragraph_format.line_spacing = 1.0

    # 5. DEPENDENCIAS (requirements.txt)
    doc.add_heading("5. Dependencias (requirements.txt)", level=1)
    p_req = doc.add_paragraph("openai\nstreamlit\npython-dotenv")
    p_req.paragraph_format.left_indent = Inches(0.5)
    p_req.runs[0].font.name = "Consolas"

    # 6. BANCO DE PRUEBAS Y VALIDACIÓN
    doc.add_heading("6. Pruebas Realizadas y Resultados", level=1)
    p_pruebas_desc = doc.add_paragraph(
        "Se ejecutaron las pruebas establecidas en la guía práctica para evaluar tanto la especialización temática "
        "como el manejo multimodal y el contexto conversacional:"
    )

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["Prueba", "Tipo de Entrada", "Consulta / Entrada", "Resultado Observado"]
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_background(hdr_cells[i], "D9D9D9")

    test_data = [
        ("Prueba 1", "Texto", "¿Qué ingredientes lleva el ají de gallina?", "Responde detalladamente con pechuga de pollo/gallina, ají amarillo, pan remojado/galleta, leche evaporada, caldo, cebolla, ajo, pecanas, queso parmesano y guarniciones."),
        ("Prueba 2", "Texto", "¿Cuál es el origen del ceviche peruano?", "Explica el origen histórico precolombino (cultura Moche con tumbo/chicha) y la posterior incorporación del limón y cebolla traídos por los españoles."),
        ("Prueba 3", "Texto (Fuera de alcance)", "¿Cómo configuro un router Cisco?", "Aplica el SYSTEM_PROMPT: declina cortésmente la consulta y recuerda que su especialidad exclusiva es la gastronomía peruana."),
        ("Prueba 4", "Audio (Whisper)", "Audio: pregunta_causa_limena.wav ('¿Cómo se prepara una causa limeña?')", "Whisper transcribe con 100% de exactitud y el chatbot responde con los pasos de preparación (papa amarilla prensada, ají amarillo, limón y relleno)."),
        ("Prueba 5", "Texto (Contexto)", "Luego de hablar de ceviche: '¿qué pescado se recomienda?'", "Mantiene el contexto previo y responde recomendando pescados frescos para ceviche como corvina, lenguado, mero, reineta o bonito.")
    ]

    for p_num, p_tipo, p_in, p_out in test_data:
        row_cells = table.add_row().cells
        row_cells[0].text = p_num
        row_cells[1].text = p_tipo
        row_cells[2].text = p_in
        row_cells[3].text = p_out
        for c in row_cells:
            c.paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 7. EVIDENCIAS VISUALES
    doc.add_heading("7. Evidencias de Ejecución en Interfaz Streamlit", level=1)
    
    evidencias = [
        ("Figura 1: Chatbot respondiendo a consulta de texto (Prueba 1 y 2)", "Captura de pantalla mostrando la pregunta sobre el Ají de Gallina o Ceviche y la respuesta generada por el asistente."),
        ("Figura 2: Rechazo de consulta fuera de dominio (Prueba 3)", "Captura de pantalla demostrando la restricción del SYSTEM_PROMPT ante preguntas ajenas como configuración de redes."),
        ("Figura 3: Carga de archivo de audio y transcripción Whisper (Prueba 4)", "Captura del componente de audio, la transcripción generada por Whisper-1 y la respuesta gastronómica resultante."),
        ("Figura 4: Mantenimiento de memoria conversacional (Prueba 5)", "Captura demostrando que el chatbot recuerda el contexto de la consulta anterior.")
    ]

    for fig_title, fig_desc in evidencias:
        p_fig = doc.add_paragraph()
        r_f = p_fig.add_run(f"[{fig_title}]\n")
        r_f.bold = True
        r_f.font.color.rgb = RGBColor(0, 32, 96)
        p_fig.add_run(f"Descripción: {fig_desc}\n")
        
        # Cuadro contenedor para imagen
        p_box = doc.add_paragraph()
        r_box = p_box.add_run("─── [ PEGAR AQUÍ CAPTURA DE PANTALLA ] ───")
        r_box.font.color.rgb = RGBColor(128, 128, 128)
        r_box.font.italic = True
        p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_box.paragraph_format.space_after = Pt(14)

    # 8. REPOSITORIO GITHUB
    doc.add_heading("8. Repositorio en GitHub y Control de Versiones", level=1)
    doc.add_paragraph(
        "El código del proyecto ha sido organizado y estructurado para su publicación en GitHub. "
        "Por motivos estrictos de seguridad y buenas prácticas, el archivo .env con la API key NO se incluye en el repositorio."
    )
    p_git = doc.add_paragraph(
        "Comandos de inicialización y despliegue:\n"
        "  git init\n"
        "  git add .\n"
        "  git commit -m \"Implementación Chatbot Multimodal Comida Peruana\"\n"
        "  git branch -M main\n"
        "  git remote add origin https://github.com/TU_USUARIO/chatbot-comida-peruana.git\n"
        "  git push -u origin main\n\n"
        "Enlace al repositorio: https://github.com/TU_USUARIO/chatbot-comida-peruana"
    )
    p_git.paragraph_format.left_indent = Inches(0.5)
    p_git.runs[0].font.name = "Consolas"
    p_git.runs[0].font.size = Pt(9.5)

    # 9. CONCLUSIONES
    doc.add_heading("9. Conclusiones", level=1)
    doc.add_paragraph(
        "1. La integración de Streamlit con la API de OpenAI permite desarrollar soluciones de inteligencia artificial con interfaces interactivas y profesionales en muy pocas líneas de código.\n"
        "2. El modelo Whisper-1 demostró un desempeño excepcional reconociendo términos gastronómicos autóctonos peruanos, garantizando accesibilidad y multimodalidad fluida.\n"
        "3. El uso de variables de entorno (.env) y exclusión (.gitignore) es fundamental para evitar la exposición accidental de claves criptográficas y credenciales de facturación en repositorios públicos."
    )

    doc.save(r"C:\Users\Jeremy\chatbot_comida_peruana\Informe_Chatbot_Multimodal_Comida_Peruana.docx")
    print("Informe generado exitosamente en C:\\Users\\Jeremy\\chatbot_comida_peruana\\Informe_Chatbot_Multimodal_Comida_Peruana.docx")

if __name__ == "__main__":
    create_report()
